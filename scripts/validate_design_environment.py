from pathlib import Path
import json, subprocess, hashlib, os, tempfile
import ezdxf, openpyxl, docx, pptx, pypdf, PIL, matplotlib
from PIL import Image, ImageDraw
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl.styles import Font
from pptx.util import Inches, Pt
from matplotlib import pyplot as plt
from matplotlib.font_manager import FontProperties
ROOT = Path(tempfile.mkdtemp(prefix='environment-check-', dir=Path.cwd() / 'work'))
os.environ['XDG_CACHE_HOME'] = str(ROOT / 'cache')
os.environ['MPLCONFIGDIR'] = str(ROOT / 'mpl')
FONT = 'Noto Sans CJK SC'
TEXT = '室内设计环境验证：中文字体、材料预算与空间方案'
synthetic = ezdxf.new('R2007'); synthetic.units = 4
synthetic.modelspace().add_lwpolyline([(0,0),(1000,0),(1000,1000),(0,1000)], close=True)
synthetic.saveas(ROOT / 'roundtrip.dxf')
d = ezdxf.readfile(ROOT / 'roundtrip.dxf')
assert d.dxfversion == 'AC1021' and d.units == 4
p = list(d.modelspace().query('LWPOLYLINE'))
assert len(p) == 1 and p[0].closed and len(p[0]) == 4
doc = docx.Document()
doc.add_heading('室内设计环境验证', 0)
r = doc.add_paragraph().add_run(TEXT)
r.font.name = FONT
r._element.get_or_add_rPr().append(OxmlElement('w:rFonts'))
r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
doc.add_paragraph('这是合成测试，单位为毫米；不代表用户项目尺寸。')
doc.save(ROOT / 'word.docx')
w = openpyxl.Workbook(); s = w.active; s.title = '材料预算'
s.append(['室内设计环境验证', '数量', '单价', '金额'])
s.append(['浅橡木', 2, 100, '=B2*C2'])
for row in s:
    for c in row: c.font = Font(name=FONT, size=12)
for col in 'ABCD': s.column_dimensions[col].width = 24
s.sheet_properties.pageSetUpPr.fitToPage = True
s.page_setup.orientation = 'landscape'; s.page_setup.fitToWidth = 1
w.save(ROOT / 'excel.xlsx')
assert openpyxl.load_workbook(ROOT / 'excel.xlsx').active['D2'].value == '=B2*C2'
pr = pptx.Presentation(); sl = pr.slides.add_slide(pr.slide_layouts[6])
tf = sl.shapes.add_textbox(Inches(0.5), Inches(1), Inches(9), Inches(3)).text_frame
tf.word_wrap = True
tf.text = TEXT
for paragraph in tf.paragraphs:
    for run in paragraph.runs: run.font.name = FONT; run.font.size = Pt(24)
pr.save(ROOT / 'slides.pptx')
fig, ax = plt.subplots(figsize=(8, 3)); ax.axis('off')
ax.text(.05, .5, TEXT, fontproperties=FontProperties(fname='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'), fontsize=14)
fig.savefig(ROOT / 'direct.pdf'); plt.close(fig)
for src in ['word.docx', 'excel.xlsx', 'slides.pptx']:
    profile = ROOT / ('profile-' + src)
    subprocess.run(['soffice', '--headless', '-env:UserInstallation=' + profile.as_uri(), '--convert-to', 'pdf', '--outdir', str(ROOT), str(ROOT / src)], check=True, timeout=90)
for name in ['word', 'excel', 'slides', 'direct']:
    pdf = ROOT / (name + '.pdf')
    reader = pypdf.PdfReader(pdf)
    assert len(reader.pages) == 1, (name, len(reader.pages))
    txt = subprocess.check_output(['pdftotext', str(pdf), '-']).decode()
    assert '室内设计环境验证' in txt.replace(' ', '').replace('\n', ''), (name, txt)
    if name == 'excel': assert '200' in txt
    subprocess.run(['pdftoppm', '-scale-to', '1400', '-png', '-singlefile', str(pdf), str(ROOT / name)], check=True, timeout=60)
    with Image.open(ROOT / (name + '.png')) as im:
        assert im.width > 500 and im.getextrema()[0][0] < 200
    print(name, 'PDF text and preview passed')
im = Image.new('RGB', (640, 480), 'white'); ImageDraw.Draw(im).rectangle((80, 80, 560, 400), outline='black', width=3)
im.save(ROOT / 'model.png'); im.save(ROOT / 'style.webp')
before = hashlib.sha256((ROOT / 'model.png').read_bytes()).hexdigest()
cmd = [str(Path.cwd() / '.venv/bin/python'), 'scripts/prepare_render.py', '--image', str(ROOT / 'model.png'), '--reference', str(ROOT / 'style.webp'), '--style', '现代简约', '--materials', '浅橡木、米白墙面', '--lighting', '柔和日光', '--out', str(ROOT / 'render-job')]
subprocess.run(cmd, check=True)
m = json.loads((ROOT / 'render-job/manifest.json').read_text())
assert m['status'] == 'prepared_not_rendered' and m['review_status'] == 'not_generated'
assert m['target']['sha256'] == before and len(m['references']) == 1
assert '浅橡木' in (ROOT / 'render-job/prompt.txt').read_text()
assert subprocess.run(cmd, capture_output=True).returncode != 0
assert hashlib.sha256((ROOT / 'model.png').read_bytes()).hexdigest() == before
print('Imports, DXF, Chinese Office/PDF previews and prepare_render checks passed; no render generated.')
